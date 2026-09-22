from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile


SUPPORTED_RESUME_TYPES = ["txt", "md", "pdf", "docx", "rtf"]


class ResumeParseError(ValueError):
    pass


def extract_resume_text(file_name: str, data: bytes) -> str:
    suffix = Path(file_name).suffix.lower().lstrip(".")
    if suffix in {"txt", "md"}:
        return _decode_text(data)
    if suffix == "pdf":
        return _extract_pdf(data)
    if suffix == "docx":
        return _extract_docx(data)
    if suffix == "rtf":
        return _extract_rtf(data)
    raise ResumeParseError(f"暂不支持 .{suffix or 'unknown'} 格式，请上传 txt、md、pdf、docx 或 rtf。")


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore").strip()


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ResumeParseError("解析 PDF 需要安装 pypdf，请先执行 python -m pip install -r requirements.txt。") from exc

    reader = PdfReader(BytesIO(data))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    text = text.strip()
    if not text:
        raise ResumeParseError("这个 PDF 没有提取到文字，可能是扫描件图片版，请先 OCR 后再上传。")
    return text


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ResumeParseError("解析 DOCX 需要安装 python-docx，请先执行 python -m pip install -r requirements.txt。") from exc

    try:
        document = Document(BytesIO(data))
    except (BadZipFile, ValueError) as exc:
        raise ResumeParseError("DOCX 文件解析失败，请确认文件没有损坏，且不是旧版 .doc 文件。") from exc
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ResumeParseError("这个 DOCX 没有提取到文字，请检查文件内容。")
    return text


def _extract_rtf(data: bytes) -> str:
    text = _decode_text(data)
    text = text.replace("\\par", "\n").replace("\\line", "\n")
    cleaned = []
    skip = False
    for char in text:
        if char == "{":
            skip = True
            continue
        if char == "}":
            skip = False
            continue
        if skip or char == "\\":
            continue
        cleaned.append(char)
    return "".join(cleaned).strip()
