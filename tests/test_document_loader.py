from io import BytesIO

import pytest

from career_assistant.document_loader import ResumeParseError, extract_resume_text


def test_extract_resume_text_from_txt_bytes():
    text = extract_resume_text("resume.txt", "张三\nPython SQL 数据分析".encode("utf-8"))

    assert "张三" in text
    assert "Python" in text


def test_extract_resume_text_from_docx_bytes():
    from docx import Document

    document = Document()
    document.add_paragraph("李四")
    document.add_paragraph("项目经历：RAG 求职助手")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_resume_text("resume.docx", buffer.getvalue())

    assert "李四" in text
    assert "RAG 求职助手" in text


def test_extract_resume_text_rejects_unknown_format():
    with pytest.raises(ResumeParseError):
        extract_resume_text("resume.xls", b"not a resume")
